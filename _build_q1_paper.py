# -*- coding: utf-8 -*-
"""
_build_q1_paper.py —— 问题一单篇论文的一键构建入口

把三段串起来：

    _gen_paper_tables_q1.py      → _tables_q1/*.md        （表格从 data/q1_v2/ 现取）
    _inject_paper_tables.py      → 问题一论文_展开.md      （占位符 → 表体）
    _make_paper_docx.py          → 问题一论文.docx

为什么要有这个入口，而不是让作者手敲三条命令：

1. **顺序有依赖，且错序是静默的。** 直接跑 `_make_paper_docx.py` 喂源稿，`{TABLE:...}`
   会原样漏进 Word——产物打开看着正常，内容里却有占位符。转换器已加了一道正则拦截，
   本脚本再保证正常路径上不会踩到它。
2. **既有交付物必须可证未被破坏。** 本脚本在动手前后各取一次 `论文.docx` 与
   `data/q1_delivery/` 的指纹，**不相等就报错退出**。三问稿与旧问题一交付是本次
   工作的「旁观者」，它们被改动只可能是失误。
3. **输出名写死。** 绝不提供 `--out 论文.docx`，也没有这个选项——覆盖三问稿是
   本脚本唯一能造成的不可逆破坏，故在参数层面就把它拿掉。

用法：
    python _build_q1_paper.py              # 全链
    python _build_q1_paper.py --stage=tables   # 只重出表格
    python _build_q1_paper.py --stage=md       # 出到展开稿为止
"""

from __future__ import annotations

import argparse
import hashlib
import os
import re
import subprocess
import sys

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")  # type: ignore[attr-defined]
    except Exception:
        pass

ROOT = os.path.dirname(os.path.abspath(__file__))
PY = sys.executable

SRC = os.path.join(ROOT, "问题一论文.md")
EXPANDED = os.path.join(ROOT, "问题一论文_展开.md")
DOCX = os.path.join(ROOT, "问题一论文.docx")
TABLES = os.path.join(ROOT, "_tables_q1")

#: 本脚本**不得**触碰的既有交付物。三问稿与旧问题一交付是旁观者。
PROTECTED = (
    os.path.join(ROOT, "论文.docx"),
    os.path.join(ROOT, "论文_正文_展开.md"),
)
PROTECTED_DIRS = (os.path.join(ROOT, "data", "q1_delivery"),)


def _file_sha(path: str) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _dir_digest(path: str) -> str:
    """目录指纹：文件名 + 大小 + mtime 的排序摘要。

    不逐字节哈希：`data/q1_delivery/` 有 27 个文件、约 15 MiB，逐字节读一遍要几秒，
    而本脚本只是要发现「被动过」这件事。文件集合、大小与 mtime 三者一起变而未察觉，
    不是失误能造成的。
    """
    if not os.path.isdir(path):
        return "<absent>"
    items = []
    for dirpath, dirnames, filenames in os.walk(path):
        dirnames[:] = sorted(d for d in dirnames if d != "__pycache__")
        for fn in sorted(filenames):
            fp = os.path.join(dirpath, fn)
            st = os.stat(fp)
            items.append(f"{os.path.relpath(fp, path)}|{st.st_size}|{int(st.st_mtime)}")
    return hashlib.sha256("\n".join(items).encode("utf-8")).hexdigest()


def _snapshot() -> dict:
    snap = {}
    for p in PROTECTED:
        snap[p] = _file_sha(p) if os.path.isfile(p) else "<absent>"
    for d in PROTECTED_DIRS:
        snap[d] = _dir_digest(d)
    return snap


def _run(cmd: list, label: str) -> None:
    print(f"\n{'=' * 72}\n[{label}] {' '.join(cmd)}\n{'=' * 72}")
    r = subprocess.run(cmd, cwd=ROOT)
    if r.returncode != 0:
        raise SystemExit(f"[失败] {label} 退出码 {r.returncode}")


def _identity_scan(paths: list) -> list:
    """对新建的根级产物做身份扫描。

    根级 `.md` / `.docx` **不在** `code/package_check.py` 的扫描范围内（它只扫交付物
    目录与 `code/`），故这里必须自己扫一遍——否则「身份扫描零命中」这句话对新论文
    根本不成立。
    """
    sys.path.insert(0, os.path.join(ROOT, "code"))
    import package_check as PC

    hits = []
    for p in paths:
        if not os.path.isfile(p):
            continue
        if p.lower().endswith(".docx"):
            try:
                from docx import Document
                txt = "\n".join(par.text for par in Document(p).paragraphs)
                for tb in Document(p).tables:
                    for row in tb.rows:
                        txt += "\n" + "\t".join(c.text for c in row.cells)
            except Exception as e:  # noqa: BLE001
                print(f"  ⚠ 无法读取 {os.path.basename(p)}：{e}")
                continue
        else:
            with open(p, "r", encoding="utf-8", errors="ignore") as f:
                txt = f.read()
        for pat, desc in PC.IDENTITY_PATTERNS:
            m = pat.search(txt)
            if m:
                hits.append((os.path.basename(p), desc, m.group(0)[:80]))
    return hits


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="问题一单篇论文构建链")
    ap.add_argument("--stage", default="all",
                    choices=("tables", "md", "docx", "all"),
                    help="tables=只出表；md=出到展开稿；docx/all=全链")
    args = ap.parse_args(argv)

    # 输出名是常量。这里再挡一次，是为了让「本脚本不会覆盖论文.docx」这句话
    # 在代码里可见，而不是只存在于注释里。
    assert os.path.basename(DOCX) == "问题一论文.docx", DOCX

    before = _snapshot()
    print("[基线] 既有交付物指纹已记录：")
    for k, v in before.items():
        print(f"       {os.path.relpath(k, ROOT):<34}{v[:16]}…")

    _run([PY, os.path.join(ROOT, "_gen_paper_tables_q1.py")], "1/3 生成表格片段")

    if args.stage == "tables":
        return _finish(before)

    _run([PY, os.path.join(ROOT, "_inject_paper_tables.py"),
          "--src", SRC, "--out", EXPANDED, "--tables", TABLES], "2/3 注入表格")

    # 展开稿里残留占位符 = 有一张表没生成。此时出 docx 会把占位符写进 Word，
    # 故在这里就拦住，而不是等后面由转换器报错。
    with open(EXPANDED, encoding="utf-8") as f:
        expanded = f.read()
    leftover = re.findall(r"\{[A-Z_]+:[A-Za-z0-9_]+\}|\{STRUCT_ABLATION\}", expanded)
    if leftover:
        raise SystemExit(f"[失败] 展开稿仍有未替换占位符：{sorted(set(leftover))}")

    if args.stage == "md":
        return _finish(before)

    _run([PY, os.path.join(ROOT, "_make_paper_docx.py"),
          "--md", EXPANDED, "--out", DOCX], "3/3 生成 docx")

    return _finish(before, scanned=(SRC, EXPANDED, DOCX))


def _finish(before: dict, scanned: tuple = ()) -> int:
    after = _snapshot()
    bad = [k for k in before if before[k] != after[k]]
    print(f"\n[旁观者核验] 论文.docx 与 data/q1_delivery/ 是否被改动")
    for k in before:
        mark = "✗ 已改动" if before[k] != after[k] else "✓ 未改动"
        print(f"       {os.path.relpath(k, ROOT):<34}{mark}")
    if bad:
        print("\n[失败] 下列既有交付物被改动，本次构建不可接受：")
        for k in bad:
            print(f"       {k}")
        return 1

    if scanned:
        print(f"\n[身份扫描] 新建的根级产物（package_check.py 不覆盖此处，须自扫）")
        hits = _identity_scan(list(scanned))
        if not hits:
            print("       零命中")
        else:
            for fn, desc, s in hits:
                print(f"       ⚠ {fn}: {desc}  （样例：{s}）")
            return 1

    print("\n[完成] 问题一单篇论文构建链走通")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
